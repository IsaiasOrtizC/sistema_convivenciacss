"""Alumno Controller: consulta y administración de alumnos y sus antecedentes."""

import logging

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.core.validators import validate_email
from django.db import connection
from django.db.models import Q
from django.shortcuts import get_object_or_404, redirect, render
from openpyxl import load_workbook

from ..models import (
    Caso,
    Curso,
    Estudiante,
    HistorialImportacionEstudiantes,
    Usuario,
    _quitar_tildes,
)
from ..trazabilidad import traza_consulta

logger = logging.getLogger(__name__)


def gestion_de_estudiantes_view(request):
    if 'usuario_id' not in request.session:
        return redirect('login')
    if request.session.get('usuario_rol') != 'Administrador':
        messages.error(request, "Acceso denegado: No tienes permisos de Administrador.")
        return redirect('dashboard')
    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        apaterno = request.POST.get('apaterno', '').strip()
        amaterno = request.POST.get('amaterno', '').strip()
        correo = request.POST.get('correo', '').strip()
        curso_id = request.POST.get('curso')
        numero_lista = request.POST.get('numero_lista', '').strip()

        try:
            if not all((nombre, apaterno, amaterno)):
                raise ValidationError("Completa los nombres y apellidos del alumno.")
            if correo:
                validate_email(correo)
            curso_instancia = Curso.objects.get(pk=curso_id)
            numero_lista_val = int(numero_lista) if numero_lista.isdigit() else None

            if numero_lista_val and _numero_lista_ya_usado(curso_instancia, numero_lista_val):
                messages.error(
                    request,
                    f"Ya existe un alumno con el N° de lista {numero_lista_val} en {curso_instancia}. "
                    f"El número de lista debe ser único dentro del mismo curso."
                )
                return redirect('gestion_de_alumnos')

            nuevo_estudiante = Estudiante(
                estudiante_nombre=nombre,
                estudiante_apaterno=apaterno,
                estudiante_amaterno=amaterno,
                estudiante_correo=correo or None,
                curso=curso_instancia,
                estudiante_numero_lista=numero_lista_val,
            )
            nuevo_estudiante.save()

            if nuevo_estudiante.estudiante_identificador:
                messages.success(
                    request,
                    f"Alumno(a) {nombre} {apaterno} registrado(a) exitosamente. "
                    f"Identificador asignado: {nuevo_estudiante.estudiante_identificador}"
                )
            else:
                messages.warning(
                    request,
                    f"Alumno(a) {nombre} {apaterno} registrado(a), pero no se pudo generar su "
                    f"identificador único (revisa que el N° de lista y el curso sean válidos)."
                )
            return redirect('gestion_de_alumnos')
                
        except ValidationError as exc:
            messages.error(request, ' '.join(exc.messages))
        except Exception:
            logger.exception("Error registrando alumno %s %s", nombre, apaterno)
            messages.error(request, "No fue posible registrar al alumno. Revisa el curso y vuelve a intentarlo.")

    from urllib.parse import urlencode

    # Filtros combinables: búsqueda por texto, curso y N° de lista
    filtro_busqueda = request.GET.get('busqueda', '').strip()
    filtro_curso = request.GET.get('curso', '').strip()
    filtro_lista = request.GET.get('numero_lista', '').strip()
    filtro_estado = request.GET.get('estado', '').strip()
    traza_consulta(
        'alumnos', busqueda=filtro_busqueda, curso=filtro_curso,
        n_lista=filtro_lista, estado=filtro_estado,
    )

    estudiantes_registrados = Estudiante.objects.select_related('curso').all().order_by(
        'curso__curso_grado',
        'curso__curso_letra',
        'estudiante_apaterno',
        'estudiante_nombre'
    )

    # Búsqueda: cada palabra escrita debe coincidir con nombre, apellidos o identificador
    if filtro_busqueda:
        for termino in filtro_busqueda.split():
            termino_norm = _quitar_tildes(termino)
            estudiantes_registrados = estudiantes_registrados.filter(
                Q(estudiante_nombre__unaccent__icontains=termino_norm) |
                Q(estudiante_apaterno__unaccent__icontains=termino_norm) |
                Q(estudiante_amaterno__unaccent__icontains=termino_norm) |
                Q(estudiante_identificador__icontains=termino)
            )

    # Curso (se ignoran valores que no sean un id numérico)
    if filtro_curso.isdigit():
        estudiantes_registrados = estudiantes_registrados.filter(curso_id=int(filtro_curso))
    else:
        filtro_curso = ''

    # N° de lista
    if filtro_lista.isdigit():
        estudiantes_registrados = estudiantes_registrados.filter(estudiante_numero_lista=int(filtro_lista))
    else:
        filtro_lista = ''

    if filtro_estado in ('activo', 'inactivo'):
        estudiantes_registrados = estudiantes_registrados.filter(
            estudiante_activo=(filtro_estado == 'activo')
        )
    else:
        filtro_estado = ''

    cursos_disponibles = Curso.objects.all().order_by('curso_grado', 'curso_letra')

    # Para que los enlaces de paginación conserven los filtros activos
    parametros = {k: val for k, val in (
        ('busqueda', filtro_busqueda), ('curso', filtro_curso),
        ('numero_lista', filtro_lista), ('estado', filtro_estado)
    ) if val}
    querystring_filtros = ('&' + urlencode(parametros)) if parametros else ''

    # Paginación: 10 alumnos por página, con flechas Anterior/Siguiente
    paginador = Paginator(estudiantes_registrados, 10)
    numero_pagina = request.GET.get('pagina')
    pagina_estudiantes = paginador.get_page(numero_pagina)

    context = {
        'estudiantes': pagina_estudiantes,  # Page object: iterable + tiene .has_next, .has_previous, etc.
        'cursos': cursos_disponibles,
        'filtros': {
            'busqueda': filtro_busqueda,
            'curso': filtro_curso,
            'numero_lista': filtro_lista,
            'estado': filtro_estado,
        },
        'hay_filtros_activos': bool(parametros),
        'querystring_filtros': querystring_filtros,
        'historial_importaciones': HistorialImportacionEstudiantes.objects.all()[:10],
    }
    return render(request, 'gestion_casos/gestion_de_estudiantes.html', context)


# ============================================================
#  RF-70: IMPORTACIÓN MASIVA DE ALUMNOS DESDE EXCEL (.XLSX)
# ============================================================

def importar_estudiantes_view(request):
    if 'usuario_id' not in request.session:
        return redirect('login')
    if request.session.get('usuario_rol') != 'Administrador':
        messages.error(request, "Acceso denegado: No tienes permisos de Administrador.")
        return redirect('dashboard')

    if request.method != 'POST' or not request.FILES.get('archivo_excel'):
        messages.error(request, "Debes seleccionar un archivo Excel (.xlsx) para importar.")
        return redirect('gestion_de_alumnos')

    archivo = request.FILES['archivo_excel']

    if not archivo.name.lower().endswith('.xlsx'):
        messages.error(request, "El archivo debe estar en formato .xlsx.")
        return redirect('gestion_de_alumnos')

    try:
        libro = load_workbook(archivo, data_only=True)
        hoja = libro.active
    except Exception as e:
        messages.error(request, f"No se pudo leer el archivo. Verifica que sea un Excel válido. Detalle: {str(e)}")
        return redirect('gestion_de_alumnos')

    creados = 0
    omitidos = 0
    errores = []
    alumnos_agregados = []
    duplicados_detectados = []

    # Fila 1 = encabezados. Los datos parten en la fila 2.
    # Columnas esperadas: A) Nombres  B) Apellido Paterno  C) Apellido Materno  D) Curso (Grado)  E) Curso (Letra)  F) N° Lista (opcional)
    for num_fila, fila in enumerate(hoja.iter_rows(min_row=2, max_col=6, values_only=True), start=2):
        if not fila or all(c is None or str(c).strip() == '' for c in fila):
            continue  # fila completamente vacía, se ignora sin marcarla como error

        try:
            nombre = str(fila[0]).strip() if fila[0] is not None else ''
            apaterno = str(fila[1]).strip() if fila[1] is not None else ''
            amaterno = str(fila[2]).strip() if fila[2] is not None else ''
            valor_curso_grado = fila[3]
            curso_letra = str(fila[4]).strip() if fila[4] is not None else ''
            valor_numero_lista = fila[5] if len(fila) > 5 else None

            if not (nombre and apaterno and amaterno and valor_curso_grado is not None and curso_letra):
                errores.append(f"Fila {num_fila}: faltan datos obligatorios.")
                omitidos += 1
                continue

            # El curso se puede escribir solo como número (1, 2, 3, 4) -> se arma "1° Medio", etc.
            # Excel suele guardar los números como celdas numéricas (int/float), no como texto.
            if isinstance(valor_curso_grado, (int, float)):
                curso_grado = f"{int(valor_curso_grado)}° Medio"
            else:
                texto_grado = str(valor_curso_grado).strip()
                curso_grado = f"{texto_grado}° Medio" if texto_grado.isdigit() else texto_grado

            curso_instancia = Curso.objects.filter(
                curso_grado__iexact=curso_grado, curso_letra__iexact=curso_letra
            ).first()

            if not curso_instancia:
                errores.append(f"Fila {num_fila}: el curso '{curso_grado} {curso_letra}' no existe en el sistema.")
                omitidos += 1
                continue

            numero_lista = None
            if isinstance(valor_numero_lista, (int, float)):
                numero_lista = int(valor_numero_lista)
            elif isinstance(valor_numero_lista, str) and valor_numero_lista.strip().isdigit():
                numero_lista = int(valor_numero_lista.strip())

            if numero_lista and _numero_lista_ya_usado(curso_instancia, numero_lista):
                errores.append(
                    f"Fila {num_fila}: el N° de lista {numero_lista} ya está en uso por otro "
                    f"alumno en {curso_instancia}. Se creó/actualizó sin número de lista asignado."
                )
                numero_lista = None  # se registra igual al alumno, pero sin el N° de lista repetido

            estudiante, creado = Estudiante.objects.get_or_create(
                estudiante_nombre=nombre,
                estudiante_apaterno=apaterno,
                estudiante_amaterno=amaterno,
                curso=curso_instancia,
                defaults={'estudiante_numero_lista': numero_lista}
            )
            if creado:
                creados += 1
                alumnos_agregados.append(
                    f"{nombre} {apaterno} {amaterno} ({curso_instancia})"
                )
            else:
                omitidos += 1
                duplicado = f"{nombre} {apaterno} {amaterno} ({curso_instancia})"
                duplicados_detectados.append(duplicado)
                errores.append(f"Fila {num_fila}: {duplicado} ya existía, se omitió.")

        except Exception as e:
            logger.exception("Error importando alumno desde fila %s del archivo %s", num_fila, archivo.name)
            errores.append(f"Fila {num_fila}: error inesperado ({str(e)}).")
            omitidos += 1

    try:
        usuario_importador = Usuario.objects.only('usuario_correo').get(
            pk=request.session['usuario_id']
        )
        HistorialImportacionEstudiantes.objects.create(
            archivo=archivo.name[:255],
            usuario_id_importo=usuario_importador.pk,
            usuario_correo_importo=usuario_importador.usuario_correo,
            alumnos_agregados=alumnos_agregados,
            duplicados_detectados=duplicados_detectados,
            filas_omitidas=omitidos,
        )
    except Exception:
        logger.exception("No fue posible registrar la auditoría de importación de alumnos")
        messages.error(
            request,
            "La importación se procesó, pero no se pudo guardar su registro de auditoría. Contacta al administrador.",
        )

    if creados > 0:
        messages.success(request, f"Importación completada: {creados} alumno(s) nuevo(s) registrado(s), {omitidos} fila(s) omitida(s).")
    else:
        messages.warning(request, f"No se registró ningún alumno nuevo. {omitidos} fila(s) omitida(s) o con datos repetidos.")

    if errores:
        detalle = " | ".join(errores[:8])
        if len(errores) > 8:
            detalle += f" ... y {len(errores) - 8} fila(s) más con problemas."
        messages.error(request, f"Detalle de filas no importadas: {detalle}")

    return redirect('gestion_de_alumnos')


def editar_estudiante_view(request, estudiante_id):
    if 'usuario_id' not in request.session:
        return redirect('login')
    if request.session.get('usuario_rol') != 'Administrador':
        messages.error(request, "Acceso denegado: No tienes permisos de Administrador.")
        return redirect('dashboard')

    estudiante = get_object_or_404(Estudiante, pk=estudiante_id)

    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        apaterno = request.POST.get('apaterno', '').strip()
        amaterno = request.POST.get('amaterno', '').strip()
        correo = request.POST.get('correo', '').strip()
        curso_id = request.POST.get('curso')
        numero_lista = request.POST.get('numero_lista', '').strip()

        try:
            if not all((nombre, apaterno, amaterno)):
                raise ValidationError("Completa los nombres y apellidos del alumno.")
            if correo:
                validate_email(correo)
            curso_instancia = Curso.objects.get(pk=curso_id)
            numero_lista_val = int(numero_lista) if numero_lista.isdigit() else None

            if numero_lista_val and _numero_lista_ya_usado(curso_instancia, numero_lista_val, excluir_estudiante_id=estudiante.estudiante_id):
                messages.error(
                    request,
                    f"Ya existe OTRO alumno con el N° de lista {numero_lista_val} en {curso_instancia}. "
                    f"El número de lista debe ser único dentro del mismo curso."
                )
                return redirect('gestion_de_alumnos')

            estudiante.estudiante_nombre = nombre
            estudiante.estudiante_apaterno = apaterno
            estudiante.estudiante_amaterno = amaterno
            estudiante.estudiante_correo = correo or None
            estudiante.curso = curso_instancia
            estudiante.estudiante_numero_lista = numero_lista_val
            estudiante.save()  # recalcula el identificador automáticamente
            messages.success(request, f"Alumno(a) {nombre} {apaterno} actualizado(a) exitosamente.")
        except ValidationError as exc:
            messages.error(request, ' '.join(exc.messages))
        except Exception:
            logger.exception("Error actualizando estudiante %s", estudiante_id)
            messages.error(request, "No fue posible actualizar al alumno. Revisa los datos ingresados.")

    return redirect('gestion_de_alumnos')


def eliminar_estudiante_view(request, estudiante_id):
    if 'usuario_id' not in request.session:
        return redirect('login')
    if request.session.get('usuario_rol') != 'Administrador':
        messages.error(request, "Acceso denegado: No tienes permisos de Administrador.")
        return redirect('dashboard')
    if request.method != 'POST':
        return redirect('gestion_de_alumnos')

    estudiante = get_object_or_404(Estudiante, pk=estudiante_id)
    estudiante.estudiante_activo = False
    estudiante.save(update_fields=['estudiante_activo'])
    messages.success(
        request,
        f"El alumno {estudiante.estudiante_nombre} {estudiante.estudiante_apaterno} fue desactivado. "
        "Sus datos y antecedentes se conservaron.",
    )

    return redirect('gestion_de_alumnos')


def cambiar_estado_estudiante_view(request, estudiante_id):
    if 'usuario_id' not in request.session:
        return redirect('login')
    if request.session.get('usuario_rol') != 'Administrador':
        messages.error(request, "Acceso denegado: No tienes permisos de Administrador.")
        return redirect('dashboard')
    if request.method != 'POST':
        return redirect('gestion_de_alumnos')

    estudiante = get_object_or_404(Estudiante, pk=estudiante_id)
    estudiante.estudiante_activo = not estudiante.estudiante_activo
    estudiante.save(update_fields=['estudiante_activo'])
    estado = 'activado' if estudiante.estudiante_activo else 'desactivado'
    messages.success(
        request,
        f"El alumno {estudiante.estudiante_nombre} {estudiante.estudiante_apaterno} fue {estado}.",
    )
    return redirect('gestion_de_alumnos')


def detalle_estudiante_view(request, estudiante_id):
    """Ficha de un alumno: sus datos + todos los casos en los que ha estado involucrado."""
    if 'usuario_id' not in request.session:
        return redirect('login')

    if request.session.get('usuario_rol') == 'Profesor Jefe':
        messages.error(request, "Acceso denegado.")
        return redirect('mis_casos_curso')

    estudiante = get_object_or_404(Estudiante.objects.select_related('curso'), pk=estudiante_id)

    # La tabla 'involucran' no tiene columna 'id', se consulta con SQL directo
    with connection.cursor() as cursor:
        cursor.execute("SELECT caso_id, tipo_falta FROM involucran WHERE estudiante_id = %s", [estudiante_id])
        filas = cursor.fetchall()

    caso_ids = [f[0] for f in filas]
    gravedad_por_caso = {f[0]: f[1] for f in filas}

    casos_qs = Caso.objects.filter(caso_id__in=caso_ids).select_related('lugar').order_by('-caso_fecha')

    casos_detalle = []
    conteo_gravedad = {'Leve': 0, 'Grave': 0, 'Gravísima': 0}

    for caso in casos_qs:
        texto = caso.caso_descripcion or ""
        titulo = texto
        if "]" in texto and "-" in texto:
            partes = texto.split("] -", 1)
            titulo = partes[0].replace("[", "").strip()

        gravedad_caso = gravedad_por_caso.get(caso.caso_id, '')
        if gravedad_caso in conteo_gravedad:
            conteo_gravedad[gravedad_caso] += 1

        casos_detalle.append({
            'caso_id': caso.caso_id,
            'titulo': titulo if titulo else f"Caso #{caso.caso_id}",
            'fecha': caso.caso_fecha,
            'lugar': caso.lugar.lugar_nombre if caso.lugar else "No especificado",
            'gravedad': gravedad_caso,
        })

    prioridad_gravedad = {'Gravísima': 0, 'Grave': 1, 'Leve': 2}
    casos_detalle.sort(
        key=lambda caso: prioridad_gravedad.get(caso['gravedad'], 3)
    )

    total_casos = len(casos_detalle)

    # ---- Estado del alumno (semáforo automático según cantidad y gravedad de casos) ----
    # Puntaje ponderado: Leve = 1, Grave = 3, Gravísima = 5.
    # Se basa SOLO en el puntaje acumulado (una sola falta gravísima aislada no dispara
    # la alerta máxima por sí sola, se necesita una acumulación real).
    puntaje = (conteo_gravedad['Leve'] * 1) + (conteo_gravedad['Grave'] * 3) + (conteo_gravedad['Gravísima'] * 5)

    if puntaje >= 10:
        estado_clave = 'alerta'
        estado_texto = 'Alerta: requiere intervención'
        estado_detalle = 'Este alumno acumula una cantidad importante de faltas graves o gravísimas. Se recomienda intervención prioritaria de Convivencia Escolar.'
    elif puntaje >= 5:
        estado_clave = 'atencion'
        estado_texto = 'Requiere atención'
        estado_detalle = 'El alumno tiene antecedentes que ameritan seguimiento cercano, aunque no son motivo de alarma todavía.'
    elif puntaje >= 1:
        estado_clave = 'seguimiento'
        estado_texto = 'Bajo seguimiento'
        estado_detalle = 'El alumno tiene antecedentes leves registrados. No requiere acción urgente, pero conviene monitorear.'
    else:
        estado_clave = 'sin_incidencias'
        estado_texto = 'Sin incidencias'
        estado_detalle = 'Este alumno no tiene casos de convivencia registrados en el sistema.'

    context = {
        'estudiante': estudiante,
        'casos': casos_detalle,
        'total_casos': total_casos,
        'conteo_gravedad': conteo_gravedad,
        'puntaje_riesgo': puntaje,
        'estado_clave': estado_clave,
        'estado_texto': estado_texto,
        'estado_detalle': estado_detalle,
    }
    return render(request, 'gestion_casos/detalle_estudiante.html', context)


def _numero_lista_ya_usado(curso, numero_lista, excluir_estudiante_id=None):
    """Verifica si otro alumno del MISMO curso ya tiene ese número de lista.
    'excluir_estudiante_id' se usa al editar, para no chocar contra sí mismo."""
    if not numero_lista or not curso:
        return False
    qs = Estudiante.objects.filter(curso=curso, estudiante_numero_lista=numero_lista)
    if excluir_estudiante_id:
        qs = qs.exclude(pk=excluir_estudiante_id)
    return qs.exists()
