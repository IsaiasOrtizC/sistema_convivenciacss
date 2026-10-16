"""Report Controller: generación y exportación de reportes (PDF y XLSX)."""

import io
from xml.sax.saxutils import escape

from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import redirect
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..permissions import tiene_acceso_estadisticas
from ..services.caso_service import filtrar_casos_del_curso
from ..services.estadistica_service import obtener_datos_estadisticas
from ..services.reporte_service import (
    crear_grafico_barras_pdf,
    crear_grafico_torta_pdf,
    datos_reporte_caso,
    valor_xlsx_seguro,
)
from .auth_casos_controller import validar_acceso_exportacion_caso, validar_profesor_jefe


def descargar_caso_pdf_view(request, caso_id):
    caso = validar_acceso_exportacion_caso(request, caso_id)
    if caso is None:
        if 'usuario_id' not in request.session:
            return redirect('login')
        return redirect('historial_casos')

    datos = datos_reporte_caso(caso)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        topMargin=32,
        bottomMargin=32,
        leftMargin=36,
        rightMargin=36,
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph('Reporte de caso de convivencia', estilos['Title']),
        Spacer(1, 10),
        Paragraph(
            f"<b>Caso:</b> #{datos['id']} · {escape(datos['titulo'])}",
            estilos['Normal'],
        ),
        Paragraph(
            f"<b>Fecha:</b> {datos['fecha'].strftime('%d/%m/%Y') if datos['fecha'] else 'Sin fecha'}"
            f" · <b>Zona:</b> {escape(datos['lugar'])}",
            estilos['Normal'],
        ),
        Paragraph(
            f"<b>Descripción:</b> {escape(datos['descripcion'])}",
            estilos['Normal'],
        ),
        Spacer(1, 14),
        Paragraph('Estudiantes involucrados', estilos['Heading2']),
    ]
    filas = [['Identificador', 'Estudiante', 'Curso', 'Tipo de falta']]
    filas.extend([
        [
            item['identificador'],
            item['nombre'],
            item['curso'],
            item['tipo_falta'],
        ]
        for item in datos['involucrados']
    ])
    if len(filas) == 1:
        filas.append(['', 'Sin estudiantes asociados', '', ''])
    tabla = Table(filas, repeatRows=1, colWidths=[95, 220, 120, 130])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#7B1A1A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D8D0C9')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F7F3ED')]),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    elementos.append(tabla)
    doc.build(elementos)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="caso_{datos["id"]}.pdf"'
    return response


def descargar_caso_xlsx_view(request, caso_id):
    caso = validar_acceso_exportacion_caso(request, caso_id)
    if caso is None:
        if 'usuario_id' not in request.session:
            return redirect('login')
        return redirect('historial_casos')

    datos = datos_reporte_caso(caso)
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Caso'
    hoja.append(['Campo', 'Detalle'])
    hoja.append(['ID', datos['id']])
    hoja.append(['Título', valor_xlsx_seguro(datos['titulo'])])
    hoja.append(['Descripción', valor_xlsx_seguro(datos['descripcion'])])
    hoja.append(['Fecha', datos['fecha']])
    hoja.append(['Zona', valor_xlsx_seguro(datos['lugar'])])
    for celda in hoja[1]:
        celda.fill = PatternFill(start_color='7B1A1A', end_color='7B1A1A', fill_type='solid')
        celda.font = Font(color='FFFFFF', bold=True)
    hoja.column_dimensions['A'].width = 20
    hoja.column_dimensions['B'].width = 72
    hoja.freeze_panes = 'A2'

    involucrados = libro.create_sheet('Estudiantes')
    involucrados.append(['Identificador', 'Estudiante', 'Curso', 'Tipo de falta'])
    for item in datos['involucrados']:
        involucrados.append([
            valor_xlsx_seguro(item['identificador']),
            valor_xlsx_seguro(item['nombre']),
            valor_xlsx_seguro(item['curso']),
            valor_xlsx_seguro(item['tipo_falta']),
        ])
    for celda in involucrados[1]:
        celda.fill = PatternFill(start_color='7B1A1A', end_color='7B1A1A', fill_type='solid')
        celda.font = Font(color='FFFFFF', bold=True)
    for columna, ancho in enumerate([18, 36, 22, 24], 1):
        involucrados.column_dimensions[get_column_letter(columna)].width = ancho
    involucrados.freeze_panes = 'A2'

    buffer = io.BytesIO()
    libro.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="caso_{datos["id"]}.xlsx"'
    return response


def estadisticas_pdf_view(request):
    if 'usuario_id' not in request.session:
        return redirect('login')
    if not tiene_acceso_estadisticas(request):
        messages.error(request, "Acceso denegado.")
        return redirect('dashboard')

    datos = obtener_datos_estadisticas(request)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=30, bottomMargin=30, leftMargin=35, rightMargin=35)
    estilos = getSampleStyleSheet()
    elementos = []

    elementos.append(Paragraph("Estadísticas de Convivencia Escolar", estilos['Title']))
    elementos.append(Paragraph(f"Periodo: {datos['filtros']['fecha_desde']} a {datos['filtros']['fecha_hasta']}", estilos['Normal']))
    elementos.append(Paragraph(f"Generado el {timezone.now().strftime('%d/%m/%Y %H:%M')}", estilos['Normal']))
    elementos.append(Spacer(1, 14))

    resumen_data = [
        ["Total Casos", "Curso c/ más incidentes", "Tipo más frecuente", "Zona c/ más incidentes", "Variación vs año anterior"],
        [str(datos['total_casos']), datos['curso_top'], datos['tipo_caso_top'], datos['zona_top'], f"{datos['variacion_pct']}%"]
    ]
    tabla_resumen = Table(resumen_data, colWidths=[85, 105, 100, 105, 95])
    tabla_resumen.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#7B1A1A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, -1), 7.5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elementos.append(tabla_resumen)
    elementos.append(Spacer(1, 18))

    elementos.append(Paragraph("Casos por Gravedad", estilos['Heading3']))
    elementos.append(crear_grafico_barras_pdf(datos['gravedad_labels'], datos['gravedad_data']))
    elementos.append(Spacer(1, 10))

    elementos.append(Paragraph("Casos por Zona del Colegio", estilos['Heading3']))
    elementos.append(crear_grafico_torta_pdf(datos['zonas_labels'], datos['zonas_data']))
    elementos.append(Spacer(1, 18))

    elementos.append(Paragraph("Detalle de Casos", estilos['Heading3']))
    tabla_data = [["ID", "Título", "Fecha", "Curso(s)", "Zona", "Gravedad"]]
    for c in datos['tabla_datos']:
        tabla_data.append([
            f"#{c['caso_id']}", c['titulo'],
            c['fecha'].strftime('%d/%m/%Y') if c['fecha'] else '',
            c['cursos'], c['zona'], c['gravedad']
        ])
    if len(tabla_data) == 1:
        elementos.append(Paragraph("No hay casos que coincidan con los filtros aplicados.", estilos['Normal']))
    else:
        tabla = Table(tabla_data, repeatRows=1, colWidths=[30, 135, 55, 100, 90, 60])
        tabla.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#7B1A1A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 7),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F7F3ED')]),
        ]))
        elementos.append(tabla)

    doc.build(elementos)
    buffer.seek(0)
    nombre_archivo = f"estadisticas_convivencia_{timezone.now().strftime('%Y%m%d')}.pdf"
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    return response


def estadisticas_xlsx_view(request):
    if 'usuario_id' not in request.session:
        return redirect('login')
    if not tiene_acceso_estadisticas(request):
        messages.error(request, "Acceso denegado.")
        return redirect('dashboard')

    datos = obtener_datos_estadisticas(request)
    wb = Workbook()

    # Hoja 1: Resumen
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    ws_resumen['A1'] = "Estadísticas de Convivencia Escolar"
    ws_resumen['A1'].font = Font(bold=True, size=14, color="7B1A1A")
    ws_resumen['A2'] = f"Periodo: {datos['filtros']['fecha_desde']} a {datos['filtros']['fecha_hasta']}"
    ws_resumen.append([])
    ws_resumen.append(["Total de Casos", "Curso con más incidentes", "Tipo Más Frecuente", "Zona con más incidentes", "Variación % vs año anterior"])
    for col in range(1, 6):
        celda = ws_resumen.cell(row=4, column=col)
        celda.fill = PatternFill(start_color="7B1A1A", end_color="7B1A1A", fill_type="solid")
        celda.font = Font(color="FFFFFF", bold=True)
    ws_resumen.append([datos['total_casos'], datos['curso_top'], datos['tipo_caso_top'], datos['zona_top'], f"{datos['variacion_pct']}%"])
    for i, ancho in enumerate([14, 26, 20, 26, 24], 1):
        ws_resumen.column_dimensions[get_column_letter(i)].width = ancho

    # Hoja 2: Gravedad (gráfico de barras)
    ws_grav = wb.create_sheet("Por Gravedad")
    ws_grav.append(["Gravedad", "Total de Casos"])
    for label, valor in zip(datos['gravedad_labels'], datos['gravedad_data']):
        ws_grav.append([label, valor])
    chart_barras = BarChart()
    chart_barras.title = "Casos por Gravedad"
    chart_barras.y_axis.title = 'Casos'
    filas = max(len(datos['gravedad_labels']), 1) + 1
    chart_barras.add_data(Reference(ws_grav, min_col=2, min_row=1, max_row=filas), titles_from_data=True)
    chart_barras.set_categories(Reference(ws_grav, min_col=1, min_row=2, max_row=filas))
    ws_grav.add_chart(chart_barras, "D2")

    # Hoja 3: Cursos (gráfico de barras)
    ws_cursos = wb.create_sheet("Por Curso")
    ws_cursos.append(["Curso", "Total de Casos"])
    for label, valor in zip(datos['cursos_labels'], datos['cursos_data']):
        ws_cursos.append([label, valor])
    chart_cursos = BarChart()
    chart_cursos.title = "Casos por Curso"
    filas = max(len(datos['cursos_labels']), 1) + 1
    chart_cursos.add_data(Reference(ws_cursos, min_col=2, min_row=1, max_row=filas), titles_from_data=True)
    chart_cursos.set_categories(Reference(ws_cursos, min_col=1, min_row=2, max_row=filas))
    ws_cursos.add_chart(chart_cursos, "D2")

    # Hoja 4: Zonas (gráfico de torta)
    ws_zonas = wb.create_sheet("Por Zona")
    ws_zonas.append(["Zona", "Total de Casos"])
    for label, valor in zip(datos['zonas_labels'], datos['zonas_data']):
        ws_zonas.append([label, valor])
    chart_torta = PieChart()
    chart_torta.title = "Casos por Zona del Colegio"
    filas = max(len(datos['zonas_labels']), 1) + 1
    chart_torta.add_data(Reference(ws_zonas, min_col=2, min_row=1, max_row=filas), titles_from_data=True)
    chart_torta.set_categories(Reference(ws_zonas, min_col=1, min_row=2, max_row=filas))
    ws_zonas.add_chart(chart_torta, "D2")

    # Hoja 5: Evolución mensual (gráfico de línea)
    ws_mensual = wb.create_sheet("Evolución Mensual")
    ws_mensual.append(["Mes", "Total de Casos"])
    for label, valor in zip(datos['meses_labels'], datos['mensual_data']):
        ws_mensual.append([label, valor])
    chart_linea = LineChart()
    chart_linea.title = "Evolución Mensual de Casos"
    filas = len(datos['meses_labels']) + 1
    chart_linea.add_data(Reference(ws_mensual, min_col=2, min_row=1, max_row=filas), titles_from_data=True)
    chart_linea.set_categories(Reference(ws_mensual, min_col=1, min_row=2, max_row=filas))
    ws_mensual.add_chart(chart_linea, "D2")

    # Hoja 6: Detalle de casos
    ws_detalle = wb.create_sheet("Detalle de Casos")
    ws_detalle.append(["ID", "Título", "Fecha", "Curso(s)", "Zona", "Gravedad"])
    for col in range(1, 7):
        celda = ws_detalle.cell(row=1, column=col)
        celda.fill = PatternFill(start_color="7B1A1A", end_color="7B1A1A", fill_type="solid")
        celda.font = Font(color="FFFFFF", bold=True)
    for c in datos['tabla_datos']:
        ws_detalle.append([
            f"#{c['caso_id']}", c['titulo'],
            c['fecha'].strftime('%d/%m/%Y') if c['fecha'] else '',
            c['cursos'], c['zona'], c['gravedad']
        ])
    for i, ancho in enumerate([8, 32, 12, 20, 20, 12], 1):
        ws_detalle.column_dimensions[get_column_letter(i)].width = ancho

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    nombre_archivo = f"estadisticas_convivencia_{timezone.now().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    return response


def descargar_casos_pdf_view(request):
    curso = validar_profesor_jefe(request)
    if not curso:
        messages.error(request, "Acceso denegado o no tienes un curso asignado como Profesor Jefe.")
        return redirect('dashboard')

    casos, _ = filtrar_casos_del_curso(request, curso)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), topMargin=30, bottomMargin=30, leftMargin=25, rightMargin=25)
    estilos = getSampleStyleSheet()
    elementos = []

    elementos.append(Paragraph(f"Bitácora de Casos de Convivencia — {curso.curso_grado} {curso.curso_letra}", estilos['Title']))
    elementos.append(Paragraph(f"Generado el {timezone.now().strftime('%d/%m/%Y %H:%M')}", estilos['Normal']))
    elementos.append(Spacer(1, 14))

    data = [["ID", "Título", "Fecha", "Lugar", "Severidad", "Alumno(s) Involucrado(s)", "Descripción"]]
    for c in casos:
        data.append([
            f"#{c['caso_id']}",
            c['titulo'],
            c['fecha'].strftime('%d/%m/%Y') if c['fecha'] else '',
            c['lugar'],
            c['tipo_falta'],
            c['alumnos'],
            c['descripcion'][:250],
        ])

    if len(data) == 1:
        elementos.append(Paragraph("No hay casos registrados para este curso.", estilos['Normal']))
    else:
        tabla = Table(data, repeatRows=1, colWidths=[35, 95, 55, 75, 60, 150, 250])
        tabla.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#7B1A1A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 7),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F7F3ED')]),
        ]))
        elementos.append(tabla)

    doc.build(elementos)
    buffer.seek(0)

    nombre_archivo = f"casos_{curso.curso_grado}_{curso.curso_letra}_{timezone.now().strftime('%Y%m%d')}.pdf".replace(' ', '_')
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    return response


def descargar_casos_xlsx_view(request):
    curso = validar_profesor_jefe(request)
    if not curso:
        messages.error(request, "Acceso denegado o no tienes un curso asignado como Profesor Jefe.")
        return redirect('dashboard')

    casos, _ = filtrar_casos_del_curso(request, curso)

    wb = Workbook()
    ws = wb.active
    ws.title = "Casos del Curso"

    ws.merge_cells('A1:G1')
    ws['A1'] = f"Bitácora de Casos de Convivencia — {curso.curso_grado} {curso.curso_letra}"
    ws['A1'].font = Font(bold=True, size=13, color="7B1A1A")

    ws.merge_cells('A2:G2')
    ws['A2'] = f"Generado el {timezone.now().strftime('%d/%m/%Y %H:%M')}"
    ws['A2'].font = Font(italic=True, size=9, color="777777")

    encabezados = ["ID", "Título", "Fecha", "Lugar", "Severidad", "Alumno(s) Involucrado(s)", "Descripción"]
    ws.append([])
    ws.append(encabezados)
    fila_encabezado = 4

    header_fill = PatternFill(start_color="7B1A1A", end_color="7B1A1A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_num in range(1, len(encabezados) + 1):
        celda = ws.cell(row=fila_encabezado, column=col_num)
        celda.fill = header_fill
        celda.font = header_font

    for c in casos:
        ws.append([
            f"#{c['caso_id']}",
            c['titulo'],
            c['fecha'].strftime('%d/%m/%Y') if c['fecha'] else '',
            c['lugar'],
            c['tipo_falta'],
            c['alumnos'],
            c['descripcion'],
        ])

    anchos = [8, 26, 12, 18, 12, 38, 55]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"casos_{curso.curso_grado}_{curso.curso_letra}_{timezone.now().strftime('%Y%m%d')}.xlsx".replace(' ', '_')
    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    return response
